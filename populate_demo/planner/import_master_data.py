"""Master-data import script - see README.md step 2 for the full explanation
of why this needs two passes and what the two workbook-specific fixups are
for.

Not run by Odoo, and not importable as a module (it uses ``env``, which
only exists when piped into a running ``odoo-bin shell``/``odev`` shell
session):

    MASTER_DATA_XLSX=/path/to/workbook.xlsx $PYTHON $ODOO_BIN shell \\
        --addons-path=$ADDONS_PATH -d <dbname> --no-http < planner/import_master_data.py

    # then restart the shell and run it again with:
    IMPORT_RESUME_AFTER=account.analytic.plan MASTER_DATA_XLSX=... $PYTHON $ODOO_BIN shell \\
        --addons-path=$ADDONS_PATH -d <dbname> --no-http < planner/import_master_data.py
"""
import os

import openpyxl

# No __file__-relative default: this script has no file of its own once
# piped into `odoo-bin shell`'s stdin (`exec()` doesn't set __file__), so
# guessing a path relative to it isn't possible - the caller must say where
# the workbook is.
XLSX_PATH = os.environ['MASTER_DATA_XLSX']
RESUME_AFTER = os.environ.get('IMPORT_RESUME_AFTER')  # e.g. 'account.analytic.plan' on the second pass

# The stock.picking.type sheet references its default source/destination
# locations by *name* rather than external ID, using an older Odoo naming
# convention that 19.5's built-in locations no longer nest under.
NAME_FIXUPS = {
    'Partner Locations/Customers': 'Customers',
    'Partner Locations/Vendors': 'Vendors',
}

# Excel sheet names are capped at 31 characters, truncating this one model
# name - map it back to the real model.
MODEL_FIXUPS = {
    'account.analytic.distribution.m': 'account.analytic.distribution.model',
}


def to_str(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # load() rejects '0.0' for an integer field
    return NAME_FIXUPS.get(value, str(value))


def cell(row, i):
    return row[i] if i < len(row) else None  # some rows are shorter than the header


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    resuming = RESUME_AFTER is None
    for sheet_name in wb.sheetnames:  # sheet order == dependency order, keep it
        if not resuming:
            resuming = (sheet_name == RESUME_AFTER)
            continue
        model_name = MODEL_FIXUPS.get(sheet_name, sheet_name)
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0]
        keep = [i for i, h in enumerate(header) if h not in (None, '(no import)')]
        fields = [header[i] for i in keep]
        data = [[to_str(cell(row, i)) for i in keep] for row in rows[1:] if cell(row, keep[0]) is not None]
        result = env[model_name].load(fields, data)  # noqa: F821 - `env` only exists inside odoo-bin shell
        if result['messages']:
            raise Exception(f"{model_name}: {result['messages']}")
        print(f"{model_name}: {len(data)} rows")
        env.cr.commit()  # noqa: F821
        if sheet_name == 'account.analytic.plan' and RESUME_AFTER is None:
            print("--- restart the shell, then re-run with IMPORT_RESUME_AFTER=account.analytic.plan ---")
            return


main()
